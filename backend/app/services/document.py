"""文档解析、清洗与切片。支持 PDF / Markdown / Word / TXT。"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.core.config import settings

logger = logging.getLogger(__name__)

CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
MULTI_SPACE_RE = re.compile(r"[ \t\u00a0\u3000]{2,}")
MULTI_NEWLINE_RE = re.compile(r"\n{3,}")
MD_IMAGE_RE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
MD_CODE_FENCE_RE = re.compile(r"```.*?```", re.S)
MD_INLINE_CODE_RE = re.compile(r"`([^`]*)`")
MD_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
HTML_TAG_RE = re.compile(r"<[^>]+>")

SENTENCE_END = "。！？；!?;\n"


@dataclass
class ParsedDocument:
    text: str
    meta: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------
# 解析
# ---------------------------------------------------------------------
def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    warnings: list[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"第 {idx} 页解析失败: {exc}")
            pages.append("")
    text = "\n\n".join(p.strip() for p in pages if p and p.strip())
    if not text:
        warnings.append("PDF 未提取到文本，可能是扫描件（需 OCR），已按空内容处理")
    return ParsedDocument(text=text, meta={"pages": len(reader.pages)}, warnings=warnings)


def parse_docx(data: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    blocks: list[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            blocks.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return ParsedDocument(text="\n\n".join(blocks), meta={"paragraphs": len(blocks)})


def parse_markdown(data: bytes) -> ParsedDocument:
    return ParsedDocument(text=data.decode("utf-8", "ignore"))


def parse_text(data: bytes) -> ParsedDocument:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk", "latin-1"):
        try:
            return ParsedDocument(text=data.decode(encoding), meta={"encoding": encoding})
        except UnicodeDecodeError:
            continue
    return ParsedDocument(text=data.decode("utf-8", "ignore"))


PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "md": parse_markdown,
    "markdown": parse_markdown,
    "txt": parse_text,
}


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext in {"md", "markdown"}:
        return "md"
    if ext in {"docx", "doc", "pdf", "txt"}:
        return "docx" if ext == "doc" else ext
    return ext or "txt"


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    file_type = detect_file_type(filename)
    parser = PARSERS.get(file_type)
    if parser is None:
        raise ValueError(f"不支持的文件格式：.{file_type}（支持 {sorted(settings.allowed_extensions)}）")
    parsed = parser(data)
    parsed.meta["file_type"] = file_type
    return parsed


# ---------------------------------------------------------------------
# 清洗
# ---------------------------------------------------------------------
def clean_text(raw: str, file_type: str = "txt") -> str:
    text = (raw or "").replace("\r\n", "\n").replace("\r", "\n")
    text = CONTROL_RE.sub("", text)

    if file_type == "md":
        text = MD_CODE_FENCE_RE.sub("\n", text)
        text = MD_IMAGE_RE.sub("", text)
        text = MD_INLINE_CODE_RE.sub(r"\1", text)
        text = MD_LINK_RE.sub(r"\1", text)

    text = HTML_TAG_RE.sub("", text)
    text = MULTI_SPACE_RE.sub(" ", text)

    lines: list[str] = []
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped:
            lines.append(stripped)
        elif lines and lines[-1] != "":
            lines.append("")
    text = "\n".join(lines)
    text = MULTI_NEWLINE_RE.sub("\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------
# 切片
# ---------------------------------------------------------------------
def _split_long_paragraph(paragraph: str, size: int) -> list[str]:
    """按句末标点把超长段落切开，尽量不破坏句子。"""
    if len(paragraph) <= size:
        return [paragraph]
    pieces: list[str] = []
    buffer = ""
    for ch in paragraph:
        buffer += ch
        if ch in SENTENCE_END and len(buffer) >= size * 0.6:
            pieces.append(buffer.strip())
            buffer = ""
    if buffer.strip():
        pieces.append(buffer.strip())

    final: list[str] = []
    for piece in pieces:
        if len(piece) <= size * 1.4:
            final.append(piece)
            continue
        for i in range(0, len(piece), size):
            final.append(piece[i : i + size])
    return [p for p in final if p]


def chunk_text(text: str, size: int | None = None, overlap: int | None = None) -> list[dict]:
    """按段落聚合切片，带重叠，保留标题上下文。"""
    size = size or settings.CHUNK_SIZE
    overlap = overlap if overlap is not None else settings.CHUNK_OVERLAP

    if not text.strip():
        return []

    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    units: list[tuple[str, str]] = []  # (heading_path, content)
    current_heading = ""
    for para in paragraphs:
        heading_match = MD_HEADING_RE.match(para)
        if heading_match:
            current_heading = heading_match.group(2).strip()[:80]
            continue
        for piece in _split_long_paragraph(para, size):
            units.append((current_heading, piece))

    chunks: list[dict] = []
    buffer = ""
    buffer_heading = ""
    for heading, piece in units:
        candidate = f"{buffer}\n{piece}".strip() if buffer else piece
        if len(candidate) <= size or not buffer:
            buffer = candidate
            buffer_heading = buffer_heading or heading
            continue

        chunks.append(
            {
                "content": buffer.strip(),
                "heading": buffer_heading,
                "char_count": len(buffer.strip()),
            }
        )
        # 重叠：取上一块尾部 overlap 个字符接到新块开头
        tail = buffer[-overlap:] if overlap > 0 else ""
        if tail and tail not in SENTENCE_END:
            cut = max(tail.rfind(ch) for ch in SENTENCE_END)
            if cut > 0:
                tail = tail[cut + 1 :]
        buffer = f"{tail}\n{piece}".strip() if tail else piece
        buffer_heading = heading

    if buffer.strip():
        chunks.append(
            {"content": buffer.strip(), "heading": buffer_heading, "char_count": len(buffer.strip())}
        )

    # 二次切分：仍超长的块强制切
    normalized: list[dict] = []
    for chunk in chunks:
        if chunk["char_count"] <= size * 1.5:
            normalized.append(chunk)
            continue
        content = chunk["content"]
        for i in range(0, len(content), size):
            piece = content[i : i + size].strip()
            if piece:
                normalized.append(
                    {"content": piece, "heading": chunk["heading"], "char_count": len(piece)}
                )

    return [c for c in normalized if len(c["content"]) >= 10]
